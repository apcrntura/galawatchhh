"""Builds the Excel and PDF files. openpyxl and reportlab are imported lazily."""
from io import BytesIO


def _header_lines(meta: dict) -> list[str]:
    return [
        f"Destination: {meta['destination_name']} ({meta['region']})",
        f"Period: {meta['start']} to {meta['end']}",
        f"Generated: {meta['generated']} (Philippine time) by {meta['generated_by']}",
    ]


def build_xlsx(meta: dict, days: list[dict], summary: dict, alerts: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="0D1B3E")

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "GalaWatch Visitor Report"
    ws["A1"].font = Font(bold=True, size=14)
    for i, line in enumerate(_header_lines(meta), start=2):
        ws.cell(row=i, column=1, value=line)
    rows = [
        ("People who entered (total)", summary["total_in"]),
        ("People who left (total)", summary["total_out"]),
        ("Highest number inside at one time", summary["peak_inside"]),
        ("Busiest day", summary["busiest_day"] or "No data"),
        ("Days with camera data", f"{summary['days_with_data']} of {summary['days_in_range']}"),
        ("Overcrowding alerts in period", len(alerts)),
    ]
    for i, (k, v) in enumerate(rows, start=6):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 20

    wd = wb.create_sheet("Daily")
    wd.append(["Date", "Entered (IN)", "Left (OUT)", "Peak inside", "Camera data"])
    for c in wd[1]:
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(horizontal="center")
    for d in days:
        wd.append([d["day"], d["total_in"], d["total_out"], d["peak_inside"], "Yes" if d["has_data"] else "No data"])
    for col, w in zip("ABCDE", (14, 16, 14, 14, 14)):
        wd.column_dimensions[col].width = w

    wa = wb.create_sheet("Alerts")
    wa.append(["Time (Philippine)", "People inside", "Limit", "Status"])
    for c in wa[1]:
        c.font, c.fill = head_font, head_fill
    for a in alerts:
        wa.append([a["time"], a["people_inside"], a["limit_value"], a["status"]])
    for col, w in zip("ABCD", (22, 16, 10, 18)):
        wa.column_dimensions[col].width = w

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(meta: dict, days: list[dict], summary: dict, alerts: list[dict]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm, title="GalaWatch Visitor Report")
    navy = colors.HexColor("#0D1B3E")

    def table(data, widths):
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ]))
        return t

    story = [Paragraph("GalaWatch Visitor Report", styles["Title"])]
    for line in _header_lines(meta):
        story.append(Paragraph(line, styles["Normal"]))
    story += [Spacer(1, 8 * mm), Paragraph("Summary", styles["Heading2"])]
    story.append(table([
        ["Measure", "Value"],
        ["People who entered (total)", str(summary["total_in"])],
        ["People who left (total)", str(summary["total_out"])],
        ["Highest number inside at one time", str(summary["peak_inside"])],
        ["Busiest day", summary["busiest_day"] or "No data"],
        ["Days with camera data", f"{summary['days_with_data']} of {summary['days_in_range']}"],
        ["Overcrowding alerts in period", str(len(alerts))],
    ], [95 * mm, 60 * mm]))

    story += [Spacer(1, 6 * mm), Paragraph("Daily counts", styles["Heading2"])]
    story.append(table(
        [["Date", "Entered (IN)", "Left (OUT)", "Peak inside"]]
        + [[d["day"], d["total_in"], d["total_out"], d["peak_inside"] if d["has_data"] else "-"] for d in days],
        [40 * mm, 38 * mm, 38 * mm, 38 * mm]))

    story += [Spacer(1, 6 * mm), Paragraph("Overcrowding alerts", styles["Heading2"])]
    if alerts:
        story.append(table(
            [["Time (Philippine)", "People inside", "Limit", "Status"]]
            + [[a["time"], a["people_inside"], a["limit_value"], a["status"]] for a in alerts],
            [50 * mm, 38 * mm, 25 * mm, 38 * mm]))
    else:
        story.append(Paragraph("No overcrowding alerts in this period.", styles["Normal"]))
    story += [Spacer(1, 8 * mm), Paragraph(
        "Counts come from the AI camera and show people crossing the entrance line. They are estimates and can drift.",
        styles["Italic"])]
    doc.build(story)
    return buf.getvalue()
